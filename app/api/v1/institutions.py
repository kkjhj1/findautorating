import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.models.institution import Institution, AgencyInstitutionMapping
from app.models.current_rating import CurrentRating
from app.models.enums import AgencyCode
from app.schemas.institution import (
    InstitutionResponse,
    SearchNamesUpdate,
    InstitutionCreate,
    InstitutionUpdate,
    ManualMappingUpdate
)

router = APIRouter(prefix="/institutions", tags=["Institutions"])

@router.get("", response_model=List[InstitutionResponse])
def get_institutions(
    industry: Optional[str] = Query(None),
    use_yn: Optional[str] = Query(None, description="Y, N, 또는 None(전체)"),
    search: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """3번 탭 상품제공기관 목록 조회 (사용/미사용 필터 지원)"""
    query = db.query(Institution)
    if use_yn and use_yn != "ALL":
        query = query.filter(Institution.use_yn == use_yn)
    if industry and industry != "ALL":
        query = query.filter(Institution.industry_type == industry)
    if search:
        query = query.filter(
            (Institution.institution_name.ilike(f"%{search}%")) |
            (Institution.original_name.ilike(f"%{search}%")) |
            (Institution.institution_id.ilike(f"%{search}%"))
        )
    return query.order_by(Institution.institution_id).all()

@router.get("/export/excel")
def export_institutions_excel(db: Session = Depends(get_db)):
    """[전문가 추천 요건 A] 상품제공기관 마스터 엑셀 내보내기"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "상품제공기관_마스터"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Pretendard", size=10, bold=True, color="FFFFFF")
    regular_font = Font(name="Pretendard", size=9)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = [
        "기관코드", "업권", "상품제공기관명", "1순위 검색어", "2순위 검색어", "3순위 검색어",
        "한기평 조사항목", "한신평 조사항목", "NICE 조사항목", "사용여부", "화면수정여부", "비고(메모)"
    ]
    ws.append(headers)

    institutions = db.query(Institution).order_by(Institution.institution_id).all()
    for inst in institutions:
        ws.append([
            inst.institution_id,
            inst.industry_type,
            inst.institution_name,
            inst.search_name_1 or "-",
            inst.search_name_2 or "-",
            inst.search_name_3 or "-",
            inst.kr_target_rating or "-",
            inst.kis_target_rating or "-",
            inst.nice_target_rating or "-",
            inst.use_yn,
            "Y" if inst.user_modified else "N",
            inst.memo or "-"
        ])

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 16

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [3, 12]:
                c.alignment = left_align
            else:
                c.alignment = center_align

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"institution_master_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/{institution_id}", response_model=InstitutionResponse)
def get_institution_detail(institution_id: str, db: Session = Depends(get_db)):
    inst = db.query(Institution).filter_by(institution_id=institution_id).first()
    if not inst:
        raise HTTPException(status_code=404, detail="기관을 찾을 수 없습니다.")
    return inst

@router.put("/{institution_id}/search-names")
def update_search_names(institution_id: str, payload: SearchNamesUpdate, db: Session = Depends(get_db)):
    """[요구사항 5] 상세 팝업에서 1/2/3순위 검색어 빠른 저장 (user_modified=True)"""
    inst = db.query(Institution).filter_by(institution_id=institution_id).first()
    if not inst:
        raise HTTPException(status_code=404, detail="기관을 찾을 수 없습니다.")

    fields_set = payload.model_fields_set

    if "search_name_1" in fields_set:
        inst.search_name_1 = (payload.search_name_1 or "").strip() or inst.institution_name
    if "search_name_2" in fields_set:
        inst.search_name_2 = (payload.search_name_2 or "").strip() or None
    if "search_name_3" in fields_set:
        inst.search_name_3 = (payload.search_name_3 or "").strip() or None

    inst.user_modified = True
    db.commit()
    db.refresh(inst)

    return {
        "status": "success",
        "message": "검색어가 성공적으로 저장되었습니다.",
        "institution_id": inst.institution_id,
        "search_name_1": inst.search_name_1,
        "search_name_2": inst.search_name_2,
        "search_name_3": inst.search_name_3
    }

@router.post("", response_model=InstitutionResponse)
def create_institution(payload: InstitutionCreate, db: Session = Depends(get_db)):
    """[요구사항 6] 신규 상품제공기관 등록 (M&A 신설 등)"""
    existing = db.query(Institution).filter_by(institution_name=payload.institution_name.strip()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"이미 등록된 기관명입니다: {payload.institution_name}")

    # ID 자동 채번 (INS_업권3글자_순번)
    prefix = payload.industry_type[:3].upper()
    count = db.query(Institution).filter(Institution.institution_id.like(f"INS_{prefix}_%")).count() + 1
    new_id = f"INS_{prefix}_{count:03d}"

    # 중복 방지 루프
    while db.query(Institution).filter_by(institution_id=new_id).first():
        count += 1
        new_id = f"INS_{prefix}_{count:03d}"

    clean_name = payload.institution_name.strip()
    s1 = payload.search_name_1.strip() if payload.search_name_1 else clean_name
    s2 = payload.search_name_2.strip() if payload.search_name_2 else None
    s3 = payload.search_name_3.strip() if payload.search_name_3 else None

    # 업권 기반 rating_type 기본값 추론
    inferred_rating_type = payload.rating_type
    if not inferred_rating_type:
        if "보험" in payload.industry_type or "INSURANCE" in payload.industry_type:
            inferred_rating_type = "INSURANCE_FINANCIAL_STRENGTH"
        elif "증권" in payload.industry_type or "SECURITIES" in payload.industry_type:
            inferred_rating_type = "SENIOR_BOND"
        else:
            inferred_rating_type = "CORPORATE_CREDIT"

    inst = Institution(
        institution_id=new_id,
        institution_name=clean_name,
        original_name=clean_name,
        industry_type=payload.industry_type,
        rating_type=inferred_rating_type,
        search_name_1=s1,
        search_name_2=s2,
        search_name_3=s3,
        kr_target_rating=payload.kr_target_rating or None,
        kis_target_rating=payload.kis_target_rating or None,
        nice_target_rating=payload.nice_target_rating or None,
        use_yn=payload.use_yn or "Y",
        user_modified=True,
        memo=payload.memo or None
    )
    db.add(inst)

    # 3대 평가사 매핑 기본 생성
    for ag_code in [AgencyCode.KR.value, AgencyCode.KIS.value, AgencyCode.NICE.value]:
        mapping = AgencyInstitutionMapping(
            institution_id=new_id,
            agency_code=ag_code,
            search_name=s1,
            alias_name=s2,
            mapping_status="AUTO_MAPPED",
            confidence=1.0
        )
        db.add(mapping)

    # CurrentRating 빈 레코드 초기화 (대시보드 즉시 노출용)
    curr = CurrentRating(
        institution_id=new_id,
        overall_status="NO_RATING",
        final_rating=None,
        prev_final_rating=None,
        rating_change_status="NEW",
        rating_change_label="신규",
        calculation_logic="신규 등록 기관 (조사 대기)"
    )
    db.add(curr)

    db.commit()
    db.refresh(inst)
    return inst

@router.put("/mapping")
def update_institution_mapping(payload: ManualMappingUpdate, db: Session = Depends(get_db)):
    """수동확인 후 매핑 정보 확정 (PRD 제12조, 28조)"""
    mapping = db.query(AgencyInstitutionMapping).filter_by(
        institution_id=payload.institution_id,
        agency_code=payload.agency_code
    ).first()

    if not mapping:
        mapping = AgencyInstitutionMapping(
            institution_id=payload.institution_id,
            agency_code=payload.agency_code,
            search_name=payload.search_name,
            alias_name=payload.alias_name,
            mapping_status=payload.status,
            confidence=1.0
        )
        db.add(mapping)
    else:
        mapping.search_name = payload.search_name
        if payload.alias_name:
            mapping.alias_name = payload.alias_name
        mapping.mapping_status = payload.status
        mapping.confidence = 1.0

    db.commit()
    return {"status": "success", "message": "매핑 정보가 업데이트되었습니다."}

@router.put("/{institution_id}", response_model=InstitutionResponse)
def update_institution(institution_id: str, payload: InstitutionUpdate, db: Session = Depends(get_db)):
    """[요구사항 6] 상품제공기관 정보 수정 (사명변경, 업권, 검색어, 사용여부 등)"""
    inst = db.query(Institution).filter_by(institution_id=institution_id).first()
    if not inst:
        raise HTTPException(status_code=404, detail="기관을 찾을 수 없습니다.")

    fields_set = payload.model_fields_set

    if "institution_name" in fields_set and payload.institution_name:
        inst.institution_name = payload.institution_name.strip()
    if "industry_type" in fields_set and payload.industry_type:
        inst.industry_type = payload.industry_type
    if "rating_type" in fields_set and payload.rating_type:
        inst.rating_type = payload.rating_type
    if "search_name_1" in fields_set:
        inst.search_name_1 = (payload.search_name_1 or "").strip() or inst.institution_name
    if "search_name_2" in fields_set:
        inst.search_name_2 = (payload.search_name_2 or "").strip() or None
    if "search_name_3" in fields_set:
        inst.search_name_3 = (payload.search_name_3 or "").strip() or None
    if "kr_target_rating" in fields_set:
        inst.kr_target_rating = (payload.kr_target_rating or "").strip() or None
    if "kis_target_rating" in fields_set:
        inst.kis_target_rating = (payload.kis_target_rating or "").strip() or None
    if "nice_target_rating" in fields_set:
        inst.nice_target_rating = (payload.nice_target_rating or "").strip() or None
    if "memo" in fields_set:
        inst.memo = (payload.memo or "").strip() or None
    if "use_yn" in fields_set and payload.use_yn:
        inst.use_yn = payload.use_yn

    inst.user_modified = True
    db.commit()
    db.refresh(inst)
    return inst

@router.delete("/{institution_id}")
def delete_institution(institution_id: str, db: Session = Depends(get_db)):
    """[요구사항 6] 상품제공기관 삭제 / 비활성화"""
    inst = db.query(Institution).filter_by(institution_id=institution_id).first()
    if not inst:
        raise HTTPException(status_code=404, detail="기관을 찾을 수 없습니다.")

    # use_yn을 'N'으로 비활성화 (이력 및 참조 데이터 보존)
    inst.use_yn = "N"
    inst.user_modified = True
    db.commit()

    return {
        "status": "success",
        "message": f"기관 '{inst.institution_name}'이(가) 비활성화(미사용) 처리되었습니다.",
        "institution_id": institution_id
    }
