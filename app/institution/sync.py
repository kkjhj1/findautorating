import os
import re
import openpyxl
from pathlib import Path
from sqlalchemy.orm import Session
from sqlalchemy import text
from app.models.institution import Institution, AgencyInstitutionMapping
from app.models.enums import IndustryType, RatingType, AgencyCode
from app.institution.search_candidates import clean_corporate_affix

# 엑셀 업권 한글 -> IndustryType Enum 매핑
INDUSTRY_MAP = {
    "은행": IndustryType.BANK.value,
    "저축은행": IndustryType.SAVINGS_BANK.value,
    "증권": IndustryType.SECURITIES.value,
    "생명보험": IndustryType.LIFE_INSURANCE.value,
    "손해보험": IndustryType.NON_LIFE_INSURANCE.value,
}

def migrate_institution_columns(db: Session):
    """SQLite 컬럼 누락 시 자동 ALTER TABLE 추가"""
    columns_to_add = [
        ("search_name_1", "VARCHAR(100)"),
        ("search_name_2", "VARCHAR(100)"),
        ("search_name_3", "VARCHAR(100)"),
        ("kr_target_rating", "VARCHAR(50)"),
        ("kis_target_rating", "VARCHAR(50)"),
        ("nice_target_rating", "VARCHAR(50)"),
        ("user_modified", "BOOLEAN DEFAULT 0"),
        ("memo", "VARCHAR(500)"),
    ]
    for col_name, col_type in columns_to_add:
        try:
            db.execute(text(f"ALTER TABLE institutions ADD COLUMN {col_name} {col_type}"))
            db.commit()
        except Exception:
            db.rollback()

def get_requirement_file_path() -> Path:
    """전산요건.xlsx 경로 탐색"""
    possible_paths = [
        Path("자료/전산요건.xlsx"),
        Path(__file__).resolve().parent.parent.parent / "자료" / "전산요건.xlsx",
        Path("C:/Users/환진/.gemini/antigravity/scratch/신용등급조회/자료/전산요건.xlsx"),
    ]
    for p in possible_paths:
        if p.exists():
            return p
    raise FileNotFoundError("전산요건.xlsx 파일을 찾을 수 없습니다.")

def sync_institutions_from_excel(db: Session) -> int:
    """
    '전산요건.xlsx'의 '조사요건' 시트를 읽어 DB 마스터 데이터 자동 동기화 (SSOT)
    """
    migrate_institution_columns(db)

    file_path = get_requirement_file_path()
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if "조사요건" not in wb.sheetnames:
        raise ValueError(f"'{file_path}'에 '조사요건' 시트가 존재하지 않습니다.")

    ws = wb["조사요건"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return 0

    header = rows[0]
    data_rows = rows[1:]

    synced_count = 0
    for idx, r in enumerate(data_rows, start=1):
        if not r or not r[1]:
            continue

        raw_ind = str(r[0]).strip() if len(r) > 0 and r[0] is not None else ""
        raw_name = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
        raw_s1 = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
        raw_s2 = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""
        raw_s3 = str(r[4]).strip() if len(r) > 4 and r[4] is not None else ""
        raw_kr = str(r[5]).strip() if len(r) > 5 and r[5] is not None else ""
        raw_kis = str(r[6]).strip() if len(r) > 6 and r[6] is not None else ""
        raw_nice = str(r[7]).strip() if len(r) > 7 and r[7] is not None else ""

        # 한국기업평가의 생명보험/손해보험사 신용등급종류를 '보험금지급능력평가'로 정규화
        if raw_kr in ["IFSR", "IFRS", "보험금지급능력"]:
            raw_kr = "보험금지급능력평가"

        industry_val = INDUSTRY_MAP.get(raw_ind, IndustryType.BANK.value)

        # RatingType 결정 (업권 표준 기반)
        if industry_val in [IndustryType.LIFE_INSURANCE.value, IndustryType.NON_LIFE_INSURANCE.value] or "보험" in raw_kr:
            r_type = RatingType.INSURANCE_FINANCIAL_STRENGTH.value
        elif industry_val == IndustryType.SECURITIES.value or "회사채" in raw_kr:
            r_type = RatingType.SENIOR_BOND.value
        else:
            r_type = RatingType.CORPORATE_CREDIT.value

        # DB 인스턴스 검색 (institution_name 기준)
        inst = db.query(Institution).filter_by(institution_name=raw_name).first()
        if not inst:
            # ID 채번: INS_업권코드_순번
            ind_prefix = industry_val[:3]
            inst_id = f"INS_{ind_prefix}_{idx:03d}"
            inst = Institution(
                institution_id=inst_id,
                institution_name=raw_name,
                original_name=raw_name,
                industry_type=industry_val,
                rating_type=r_type,
                search_name_1=raw_s1 or raw_name,
                search_name_2=raw_s2 or None,
                search_name_3=raw_s3 or None,
                kr_target_rating=raw_kr,
                kis_target_rating=raw_kis,
                nice_target_rating=raw_nice,
                use_yn="Y"
            )
            db.add(inst)
        else:
            inst.industry_type = industry_val
            inst.rating_type = r_type
            inst.search_name_1 = raw_s1 or raw_name
            inst.search_name_2 = raw_s2 or None
            inst.search_name_3 = raw_s3 or None
            inst.kr_target_rating = raw_kr
            inst.kis_target_rating = raw_kis
            inst.nice_target_rating = raw_nice

        # 평가사별 매핑 갱신 (우선검색어 기준)
        for a_code in [AgencyCode.KR.value, AgencyCode.KIS.value, AgencyCode.NICE.value]:
            mapping = db.query(AgencyInstitutionMapping).filter_by(
                institution_id=inst.institution_id,
                agency_code=a_code
            ).first()
            if not mapping:
                mapping = AgencyInstitutionMapping(
                    institution_id=inst.institution_id,
                    agency_code=a_code,
                    search_name=raw_s1 or raw_name,
                    alias_name=raw_s2 or None
                )
                db.add(mapping)
            else:
                mapping.search_name = raw_s1 or raw_name
                mapping.alias_name = raw_s2 or None

        synced_count += 1

    db.commit()
    return synced_count
